"""Записи, пережившие правку каталога, и граничные случаи, не покрытые другими файлами."""

import dataclasses
import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from booking.export import build_xlsx, rows
from booking.flow import BookingBot

OWNER = ("tg", "1")
CLIENT = ("tg", "5")
NOW = datetime(2026, 10, 14, 8, 0)


def book(storage, service_id=1, master_id=1, start=datetime(2026, 10, 16, 12, 0), user=CLIENT):
    return storage.book(
        platform=user[0],
        user_id=user[1],
        client_name="Мария",
        phone="+7 900 111-22-33",
        service_id=service_id,
        master_id=master_id,
        start=start,
        end=start.replace(hour=start.hour + 1),
        created_at=NOW,
    )


def without(catalog, *, master_id=None, service_id=None):
    """Каталог, из которого владелец убрал мастера или услугу."""
    return dataclasses.replace(
        catalog,
        masters={k: v for k, v in catalog.masters.items() if k != master_id},
        services={k: v for k, v in catalog.services.items() if k != service_id},
    )


@pytest.fixture
def stale_bot(catalog, storage, clock, outbox):
    """Бот после правки каталога: Анны (мастер 1) и окрашивания (услуга 3) больше нет."""
    stale = without(without(catalog, master_id=1), service_id=3)
    core = BookingBot(stale, storage, admins={"tg:1"}, clock=clock)
    core.senders["tg"] = outbox.sender("tg")
    return core


async def test_my_bookings_survive_removed_master(stale_bot, storage):
    book(storage, master_id=1)
    (reply,) = await stale_bot.handle_text(*CLIENT, "/my")
    assert "мастер №1 (нет в каталоге)" in reply.text


async def test_cancel_works_for_removed_master_and_owner_is_notified(stale_bot, storage, outbox):
    appt = book(storage, master_id=1)
    (ask,) = await stale_bot.handle_callback(*CLIENT, f"c:{appt.id}")
    assert "Отменить запись?" in ask.text
    (done,) = await stale_bot.handle_callback(*CLIENT, f"cy:{appt.id}")
    assert "Запись отменена" in done.text
    assert storage.get_appointment(appt.id).status == "cancelled"
    assert any("Клиент отменил" in r.text for _, _, r in outbox.sent)


async def test_owner_export_and_today_survive_removed_service(stale_bot, storage, clock):
    book(storage, service_id=3, master_id=2, start=datetime(2026, 10, 14, 15, 0))
    (today,) = await stale_bot.handle_text(*OWNER, "/today")
    assert "услуга №3 (нет в каталоге)" in today.text
    (export,) = await stale_bot.handle_text(*OWNER, "/export")
    sheet = load_workbook(io.BytesIO(export.file_bytes)).active
    assert sheet.cell(row=2, column=4).value == "услуга №3 (нет в каталоге)"
    assert sheet.cell(row=2, column=8).value is None


async def test_reminder_for_removed_master_still_sent(stale_bot, storage, clock, outbox):
    book(storage, master_id=1)
    clock.now = datetime(2026, 10, 15, 13, 0)
    assert await stale_bot.send_due_reminders() == 1
    assert "мастер №1" in outbox.sent[0][2].text


def test_rows_keep_price_for_known_service(catalog, storage):
    book(storage)
    (row,) = rows(catalog, storage.all_appointments())
    assert (row[3], row[4], row[7]) == ("Женская стрижка", "Анна", 1800)
    assert build_xlsx(catalog, storage.all_appointments())


async def test_menu_button_resets_input_step(bot, storage):
    await bot.handle_callback("max", "42", "t:1:2:202610171200")
    assert storage.get_client("max", "42").state == "ask_name"
    (menu,) = await bot.handle_callback("max", "42", "menu")
    assert "Здравствуйте" in menu.text
    assert storage.get_client("max", "42").state is None


async def test_reminder_skipped_for_platform_without_sender(bot, storage, clock, outbox):
    book(storage, user=("vk", "9"))
    clock.now = datetime(2026, 10, 15, 13, 0)
    assert await bot.send_due_reminders() == 0
    assert outbox.sent == []
    # помечено, чтобы не перебирать каждую минуту
    assert storage.all_appointments()[0].reminded is True


async def test_empty_and_slash_only_text(bot):
    (reply,) = await bot.handle_text(*CLIENT, "   ")
    assert "Не понял" in reply.text
    (reply,) = await bot.handle_text(*CLIENT, "/")
    assert "Не понял" in reply.text


async def test_master_without_service_in_button_returns_menu(bot):
    # Дмитрий (3) не делает женскую стрижку (1) — подделанная кнопка ведёт в меню
    for data in ("m:1:3", "d:1:3:20261015", "t:1:3:202610151200"):
        (reply,) = await bot.handle_callback(*CLIENT, data)
        assert "Здравствуйте" in reply.text, data


def test_failed_insert_rolls_back_and_storage_stays_usable(storage):
    with pytest.raises(Exception):  # noqa: B017 — нужен любой сбой вставки
        storage.book(
            platform=None,  # NOT NULL — вставка упадёт внутри транзакции
            user_id="5",
            client_name="Мария",
            phone="+7 900 111-22-33",
            service_id=1,
            master_id=1,
            start=datetime(2026, 10, 16, 12, 0),
            end=datetime(2026, 10, 16, 13, 0),
            created_at=NOW,
        )
    assert not storage.conn.in_transaction
    assert book(storage) is not None
