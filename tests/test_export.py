"""Строки выгрузки и файл Excel."""

import io
from datetime import date, datetime

from openpyxl import load_workbook

from booking.export import COLUMNS, build_xlsx, rows


def book(storage, platform, start, master_id=1, service_id=1):
    return storage.book(
        platform=platform,
        user_id="5",
        client_name="Мария",
        phone="+7 900 111-22-33",
        service_id=service_id,
        master_id=master_id,
        start=start,
        end=start.replace(hour=start.hour + 1),
        created_at=datetime(2026, 10, 14, 8, 5),
    )


def test_rows(catalog, storage):
    book(storage, "tg", datetime(2026, 10, 15, 12))
    cancelled = book(storage, "max", datetime(2026, 10, 16, 9), master_id=4, service_id=4)
    storage.cancel(cancelled.id)
    book(storage, "vk", datetime(2026, 10, 17, 9))  # неизвестный мессенджер — как есть

    first, second, third = rows(catalog, storage.all_appointments())
    assert first == [
        1,
        date(2026, 10, 15),
        "12:00",
        "Женская стрижка",
        "Анна",
        "Мария",
        "+7 900 111-22-33",
        1800,
        "Telegram",
        "активна",
        "14.10.2026 08:05",
    ]
    assert len(first) == len(COLUMNS)
    assert (second[0], second[4], second[8], second[9]) == (2, "Ольга", "MAX", "отменена")
    assert third[8] == "vk"


def test_xlsx(catalog, storage):
    book(storage, "tg", datetime(2026, 10, 15, 12))
    ws = load_workbook(io.BytesIO(build_xlsx(catalog, storage.all_appointments()))).active
    assert ws.title == "Записи"
    assert [c.value for c in ws[1]] == [t for t, _ in COLUMNS]
    assert ws.freeze_panes == "A2"
    assert ws["B2"].value.date() == date(2026, 10, 15)
    assert ws["B2"].number_format == "DD.MM.YYYY"
    assert ws["H2"].value == 1800
    assert ws.max_row == 2


def test_xlsx_empty(catalog):
    ws = load_workbook(io.BytesIO(build_xlsx(catalog, []))).active
    assert ws.max_row == 1
