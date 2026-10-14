import io
from datetime import datetime

from openpyxl import load_workbook

CLIENT = ("max", "42")


def datas(reply):
    return [b.data for row in reply.buttons for b in row]


async def press(bot, data, who=CLIENT):
    (reply,) = await bot.handle_callback(*who, data)
    return reply


async def say(bot, text, who=CLIENT):
    (reply,) = await bot.handle_text(*who, text)
    return reply


async def book_haircut(bot, start="202610151200", who=CLIENT, name="Мария", phone="+7 900 111-22-33"):
    """Женская стрижка у Анны: полный путь новым клиентом до экрана проверки."""
    await press(bot, "book", who)
    await press(bot, "s:1", who)
    await press(bot, "m:1:1", who)
    await press(bot, f"d:1:1:{start[:8]}", who)
    reply = await press(bot, f"t:1:1:{start}", who)
    if "Как к вам обращаться" in reply.text:
        await say(bot, name, who)
        reply = await say(bot, phone, who)
    return reply


async def test_start_shows_menu_and_admin_buttons(bot):
    client_menu = await say(bot, "/start")
    assert "Лаванда" in client_menu.text
    assert "x" not in datas(client_menu)
    admin_menu = await say(bot, "/start", ("tg", "1"))
    assert {"x", "today"} <= set(datas(admin_menu))


async def test_full_booking_path(bot, storage, outbox):
    services = await press(bot, "book")
    assert "s:1" in datas(services)
    masters = await press(bot, "s:1")
    assert {"m:1:1", "m:1:2", "m:1:0"} <= set(datas(masters))
    days = await press(bot, "m:1:1")
    assert "d:1:1:20261015" in datas(days)
    times = await press(bot, "d:1:1:20261015")
    assert "t:1:1:202610151200" in datas(times)

    ask_name = await press(bot, "t:1:1:202610151200")
    assert "Как к вам обращаться" in ask_name.text
    ask_phone = await say(bot, "Мария")
    assert "Телефон" in ask_phone.text
    retry = await say(bot, "не скажу")
    assert "Не похоже на номер" in retry.text
    summary = await say(bot, "8 900 111 22 33")
    assert "Чт, 15 октября, 12:00" in summary.text
    assert "1 800 ₽" in summary.text
    assert "+7 900 111-22-33" in summary.text

    done = await press(bot, "ok")
    assert "вы записаны" in done.text
    (appt,) = storage.all_appointments()
    assert (appt.client_name, appt.master_id, appt.start) == ("Мария", 1, datetime(2026, 10, 15, 12))
    # владелец в Telegram узнал о записи из MAX
    assert [(p, u) for p, u, _ in outbox.sent] == [("tg", "1")]
    assert "Новая запись" in outbox.sent[0][2].text

    # занятое время больше не предлагается
    times = await press(bot, "d:1:1:20261015", ("tg", "7"))
    assert "t:1:1:202610151200" not in datas(times)


async def test_returning_client_skips_name_and_phone(bot):
    await book_haircut(bot)
    await press(bot, "ok")
    summary = await press(bot, "t:1:1:202610161200")
    assert "Проверьте запись" in summary.text
    assert "Мария" in summary.text


async def test_same_slot_cannot_be_booked_twice(bot, storage):
    other = ("tg", "7")
    await book_haircut(bot)
    await book_haircut(bot, who=other, name="Ирина", phone="+7 900 222-33-44")
    await press(bot, "ok")
    lost = await press(bot, "ok", other)
    assert "заняли" in lost.text
    assert len(storage.all_appointments()) == 1


async def test_any_master_assigns_free_one(bot, storage):
    await press(bot, "t:1:0:202610161900")
    await say(bot, "Мария")
    summary = await say(bot, "+79001112233")
    assert "Марина" in summary.text
    await press(bot, "ok")
    assert storage.all_appointments()[0].master_id == 2


async def test_cancel_own_booking(bot, storage, outbox):
    await book_haircut(bot)
    await press(bot, "ok")
    my = await press(bot, "my")
    (appt,) = storage.all_appointments()
    assert f"c:{appt.id}" in datas(my)

    stranger = await press(bot, f"cy:{appt.id}", ("tg", "7"))
    assert "уже отменена" in stranger.text
    assert storage.get_appointment(appt.id).status == "active"

    await press(bot, f"c:{appt.id}")
    done = await press(bot, f"cy:{appt.id}")
    assert "отменена" in done.text
    assert storage.get_appointment(appt.id).status == "cancelled"
    assert "отменил" in outbox.sent[-1][2].text
    assert "Активных записей нет" in (await press(bot, "my")).text


async def test_reminder_sent_once_a_day_before(bot, clock, outbox):
    await press(bot, "t:1:2:202610171200")  # суббота, Марина
    await say(bot, "Мария")
    await say(bot, "+79001112233")
    await press(bot, "ok")
    outbox.sent.clear()

    clock.now = datetime(2026, 10, 16, 11, 0)  # за 25 часов — рано
    assert await bot.send_due_reminders() == 0
    clock.now = datetime(2026, 10, 16, 12, 0)
    assert await bot.send_due_reminders() == 1
    platform, user_id, reply = outbox.sent[0]
    assert (platform, user_id) == CLIENT
    assert "завтра в 12:00" in reply.text
    assert await bot.send_due_reminders() == 0


async def test_no_reminder_when_booked_within_a_day(bot, clock, outbox):
    await book_haircut(bot, start="202610141400")
    await press(bot, "ok")
    outbox.sent.clear()
    clock.now = datetime(2026, 10, 14, 9, 0)
    assert await bot.send_due_reminders() == 0
    assert outbox.sent == []


async def test_export_for_owner_only(bot):
    await book_haircut(bot)
    await press(bot, "ok")

    assert (await press(bot, "x")).file_bytes is None
    export = await press(bot, "x", ("tg", "1"))
    assert export.file_name == "zapisi_2026-10-14.xlsx"
    ws = load_workbook(io.BytesIO(export.file_bytes)).active
    assert ws["A1"].value == "№"
    assert [c.value for c in ws[2]][3:7] == ["Женская стрижка", "Анна", "Мария", "+7 900 111-22-33"]

    today = await say(bot, "/today", ("tg", "1"))
    assert "На сегодня записей нет" in today.text


async def test_garbage_callback_returns_menu(bot):
    for data in ("t:1:1:bad", "s:999", "m:2:1", "cy:abc", "zzz"):
        reply = await press(bot, data)
        assert "book" in datas(reply)


async def test_user_text_is_escaped(bot):
    await book_haircut(bot, name="<b>Хакер</b>")
    summary = await press(bot, "t:1:1:202610151200")
    assert "<b>Хакер</b>" not in summary.text
    assert "&lt;b&gt;Хакер&lt;/b&gt;" in summary.text
