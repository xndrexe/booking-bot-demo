"""Выгрузка записей для владельца: файл Excel и строки для Google Таблицы."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from booking.catalog import Catalog
from booking.db import ACTIVE, Appointment

PLATFORMS = {"tg": "Telegram", "max": "MAX"}
COLUMNS = [
    ("№", 6),
    ("Дата", 12),
    ("Время", 8),
    ("Услуга", 32),
    ("Мастер", 12),
    ("Клиент", 18),
    ("Телефон", 18),
    ("Цена, ₽", 10),
    ("Мессенджер", 12),
    ("Статус", 11),
    ("Создана", 17),
]


def rows(catalog: Catalog, appointments: list[Appointment]) -> list[list]:
    """Строки таблицы записей без шапки; дата — объектом date."""
    out = []
    for number, appt in enumerate(appointments, start=1):
        service = catalog.services[appt.service_id]
        out.append(
            [
                number,
                appt.start.date(),
                appt.start.strftime("%H:%M"),
                service.name,
                catalog.masters[appt.master_id].name,
                appt.client_name,
                appt.phone,
                service.price,
                PLATFORMS.get(appt.platform, appt.platform),
                "активна" if appt.status == ACTIVE else "отменена",
                appt.created_at.strftime("%d.%m.%Y %H:%M"),
            ]
        )
    return out


def build_xlsx(catalog: Catalog, appointments: list[Appointment]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Записи"
    ws.append([title for title, _ in COLUMNS])
    for idx, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7B5EA7")
        cell.alignment = Alignment(horizontal="center")

    for row in rows(catalog, appointments):
        ws.append(row)
        for cell in ws[ws.max_row]:
            # openpyxl делает формулой любую строку с «=»; имя клиента — чужой ввод, пусть будет текстом
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
        ws.cell(row=ws.max_row, column=2).number_format = "DD.MM.YYYY"
        ws.cell(row=ws.max_row, column=8).number_format = "#,##0"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
