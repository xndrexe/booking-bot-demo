"""Прогоняет настоящий диалог бота и рисует его в оформлении чата Telegram.

Тексты и кнопки — ровно те, что отдаёт ядро бота; вокруг — только рамка
мессенджера. Итог — HTML в build/screens/, снимки делает tools/shoot.py.

    .venv/bin/python tools/render_chat.py
"""

from __future__ import annotations

import asyncio
import io
import sys
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from html import escape
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from booking.catalog import load_catalog  # noqa: E402
from booking.db import Storage  # noqa: E402
from booking.flow import BookingBot, Reply  # noqa: E402

OUT = ROOT / "build" / "screens"
CLIENT = ("tg", "100")
OWNER = ("tg", "1")
# выдуманные клиенты, чтобы у мастеров было занятое время
FAKE_CLIENTS = [
    ("Елена", "+7 900 000-00-11", 1, 1, "2026-10-15T10:00"),
    ("Светлана", "+7 900 000-00-12", 3, 1, "2026-10-15T14:00"),
    ("Игорь", "+7 900 000-00-13", 2, 3, "2026-10-15T11:00"),
    ("Ольга", "+7 900 000-00-14", 4, 4, "2026-10-15T12:00"),
    ("Наталья", "+7 900 000-00-15", 1, 2, "2026-10-16T15:00"),
    ("Павел", "+7 900 000-00-16", 2, 3, "2026-10-16T18:00"),
]


@dataclass
class Bubble:
    side: str  # "in" — бот, "out" — клиент
    text: str
    time: str
    buttons: list = field(default_factory=list)
    pressed: str | None = None
    file_name: str | None = None


class Recorder:
    def __init__(self, core: BookingBot, who: tuple[str, str], clock):
        self.core, self.who, self.clock = core, who, clock
        self.bubbles: list[Bubble] = []

    def _stamp(self) -> str:
        self.clock.tick()
        return f"{self.clock():%H:%M}"

    def _add(self, replies: list[Reply]) -> None:
        for r in replies:
            self.bubbles.append(Bubble("in", r.text, self._stamp(), r.buttons, file_name=r.file_name))

    async def say(self, text: str) -> None:
        self.bubbles.append(Bubble("out", escape(text), self._stamp()))
        self._add(await self.core.handle_text(*self.who, text))

    async def press(self, data: str) -> None:
        # нажатую кнопку подсвечиваем на последнем сообщении, где она есть
        for bubble in reversed(self.bubbles):
            if any(b.data == data for row in bubble.buttons for b in row):
                bubble.pressed = data
                break
        else:
            raise ValueError(f"no button {data}")
        self._add(await self.core.handle_callback(*self.who, data))


class Clock:
    def __init__(self, start: datetime):
        self.now = start

    def tick(self) -> None:
        self.now += timedelta(seconds=20)

    def __call__(self) -> datetime:
        return self.now


CSS = """
* { box-sizing: border-box; margin: 0; }
body { font: 15px/1.38 Inter, 'Noto Color Emoji', sans-serif; background: #fff; }
.phone { width: 390px; background: linear-gradient(160deg, #b7d1a3 0%, #cfdca0 45%, #8fb7a5 100%); }
.bar { position: sticky; top: 0; display: flex; align-items: center; gap: 10px; padding: 10px 14px;
  background: #fff; border-bottom: 1px solid #e3e3e3; }
.ava { width: 40px; height: 40px; border-radius: 50%; display: grid; place-items: center; font-size: 20px;
  background: linear-gradient(135deg, #a98bd6, #7b5ea7); color: #fff; }
.bar b { font-size: 16px; display: block; }
.bar span { color: #8a8a8a; font-size: 13px; }
.caption { padding: 8px 14px; background: #7b5ea7; color: #fff; font-size: 13px; font-weight: 600; }
.chat { padding: 10px 8px 14px; display: flex; flex-direction: column; gap: 6px; }
.msg { max-width: 86%; padding: 7px 10px 6px; border-radius: 16px; position: relative; white-space: pre-wrap;
  box-shadow: 0 1px 1px rgba(0,0,0,.12); word-wrap: break-word; }
.in { background: #fff; align-self: flex-start; border-bottom-left-radius: 5px; }
.out { background: #e1fec6; align-self: flex-end; border-bottom-right-radius: 5px; }
.t { float: right; margin: 6px 0 -4px 10px; font-size: 11px; color: #9aa59a; }
.out .t { color: #5fae58; }
.kb { align-self: flex-start; width: 86%; display: flex; flex-direction: column; gap: 4px; margin-top: -2px; }
.row { display: flex; gap: 4px; }
.btn { flex: 1; text-align: center; padding: 8px 6px; border-radius: 10px; font-size: 14px; font-weight: 500;
  color: #fff; background: rgba(40, 60, 40, .32); }
.btn.on { background: rgba(123, 94, 167, .85); box-shadow: 0 0 0 2px #fff inset; }
.file { display: flex; gap: 10px; align-items: center; margin-bottom: 6px; }
.file i { width: 42px; height: 42px; border-radius: 50%; background: #1d8a4a; color: #fff; display: grid;
  place-items: center; font-style: normal; font-size: 12px; font-weight: 700; }
.file small { display: block; color: #8a8a8a; }
table { border-collapse: collapse; font: 12px/1.3 Inter, sans-serif; margin: 12px; }
td, th { border: 1px solid #d4d4d4; padding: 5px 8px; white-space: nowrap; }
th { background: #7b5ea7; color: #fff; font-weight: 600; }
tr:nth-child(odd) td { background: #f7f5fb; }
.sheet { padding: 0 0 4px; background: #fff; display: inline-block; }
.sheet h1 { font: 600 14px Inter, sans-serif; padding: 10px 12px 0; color: #333; }
"""


def render_chat(title: str, bubbles: list[Bubble]) -> str:
    parts = []
    for b in bubbles:
        file_part = ""
        if b.file_name:
            file_part = (
                f'<div class="file"><i>XLSX</i><div><b>{b.file_name}</b><small>Excel</small></div></div>'
            )
        parts.append(f'<div class="msg {b.side}">{file_part}{b.text}<span class="t">{b.time}</span></div>')
        if b.buttons:
            rows = "".join(
                '<div class="row">'
                + "".join(
                    f'<div class="btn{" on" if btn.data == b.pressed else ""}">{escape(btn.text)}</div>'
                    for btn in row
                )
                + "</div>"
                for row in b.buttons
            )
            parts.append(f'<div class="kb">{rows}</div>')
    return (
        f"<!doctype html><meta charset=utf-8><style>{CSS}</style>"
        '<div class="phone"><div class="bar"><div class="ava">✂️</div>'
        "<div><b>Лаванда · запись</b><span>бот</span></div></div>"
        f'<div class="caption">{escape(title)}</div>'
        f'<div class="chat">{"".join(parts)}</div></div>'
    )


def render_sheet(xlsx: bytes) -> str:
    ws = load_workbook(io.BytesIO(xlsx)).active
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        tag = "th" if i == 0 else "td"
        cells = [
            v.strftime("%d.%m.%Y")
            if isinstance(v, datetime)
            else f"{v:,}".replace(",", " ")
            if isinstance(v, int) and v > 999
            else v
            for v in row
        ]
        rows.append("<tr>" + "".join(f"<{tag}>{escape(str(v))}</{tag}>" for v in cells) + "</tr>")
    return (
        f"<!doctype html><meta charset=utf-8><style>{CSS}</style>"
        f'<div class="sheet"><h1>Выгрузка для владельца — лист «{ws.title}»</h1>'
        f"<table>{''.join(rows)}</table></div>"
    )


async def main() -> None:
    catalog = load_catalog(ROOT / "data" / "catalog.json")
    storage = Storage(":memory:")
    clock = Clock(datetime(2026, 10, 14, 11, 2))  # среда
    core = BookingBot(catalog, storage, admins={":".join(OWNER)}, clock=clock)
    outbox: list[tuple[str, Reply, str]] = []  # кому, что, во сколько

    async def send(user_id: str, reply: Reply) -> None:
        outbox.append((user_id, reply, f"{clock():%H:%M}"))

    core.senders["tg"] = send
    for name, phone, service_id, master_id, start in FAKE_CLIENTS:
        begin = datetime.fromisoformat(start)
        storage.book(
            platform="max" if master_id % 2 else "tg",
            user_id=f"9{master_id}",
            client_name=name,
            phone=phone,
            service_id=service_id,
            master_id=master_id,
            start=begin,
            end=begin + timedelta(minutes=catalog.services[service_id].duration_min),
            created_at=datetime(2026, 10, 12, 18, 30),
        )

    # 1. выбор услуги и мастера
    first = Recorder(core, CLIENT, clock)
    await first.say("/start")
    await first.press("book")
    await first.press("s:1")
    await first.press("m:1:1")

    # 2. день, время, имя и телефон, подтверждение
    second = Recorder(core, CLIENT, clock)
    second.bubbles = [first.bubbles[-1]]
    await second.press("d:1:1:20261015")
    await second.press("t:1:1:202610151200")
    await second.say("Мария")
    await second.say("8 900 123 45 67")
    await second.press("ok")

    # 3. напоминание за сутки, мои записи, отмена
    clock.now = datetime(2026, 10, 14, 12, 0)
    # запись на 15.10 12:00 сделана раньше чем за сутки — напоминание уходит 14.10 в 12:00
    sent_before = len(outbox)
    third = Recorder(core, CLIENT, clock)
    await core.send_due_reminders()
    (_, reminder, _), *_ = [o for o in outbox[sent_before:] if o[0] == CLIENT[1]]
    third.bubbles.append(Bubble("in", reminder.text, f"{clock():%H:%M}", reminder.buttons))
    await third.press(third.bubbles[-1].buttons[0][1].data)
    await third.press(third.bubbles[-1].buttons[0][0].data)

    # 4. владелец: уведомления и выгрузка
    owner = Recorder(core, OWNER, clock)
    for user_id, reply, sent_at in outbox:
        if user_id == OWNER[1]:
            owner.bubbles.append(Bubble("in", reply.text, sent_at))
    await owner.say("/start")
    await owner.press("x")

    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "1-usluga-i-master.html").write_text(
        render_chat("Выбор услуги и мастера", first.bubbles), encoding="utf-8"
    )
    (OUT / "2-vremya-i-zapis.html").write_text(
        render_chat("Свободное время и запись", second.bubbles), encoding="utf-8"
    )
    (OUT / "3-napominanie-i-otmena.html").write_text(
        render_chat("Напоминание за сутки и отмена", third.bubbles), encoding="utf-8"
    )
    (OUT / "4-vladelets.html").write_text(
        render_chat("Владелец: уведомления и выгрузка", owner.bubbles), encoding="utf-8"
    )
    export = await core.handle_callback(*OWNER, "x")
    (OUT / "5-tablitsa.html").write_text(render_sheet(export[0].file_bytes), encoding="utf-8")
    print(f"html: {OUT}")


if __name__ == "__main__":
    asyncio.run(main())
