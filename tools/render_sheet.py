"""Собирает выгрузку Excel на выдуманных записях и рисует её лист таблицей.

Файл — ровно тот, что бот отдаёт владельцу по кнопке «Выгрузка в Excel».
Итог — HTML в build/screens/, снимок делает tools/shoot.py.

    .venv/bin/python tools/render_sheet.py
"""

from __future__ import annotations

import asyncio
import io
import sys
from datetime import datetime, timedelta
from html import escape
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from booking.catalog import load_catalog  # noqa: E402
from booking.db import Storage  # noqa: E402
from booking.flow import BookingBot  # noqa: E402

OUT = ROOT / "build" / "screens"
OWNER = ("tg", "1")
# выдуманные записи для выгрузки
FAKE_CLIENTS = [
    ("Елена", "+7 900 000-00-11", 1, 1, "2026-10-15T10:00"),
    ("Светлана", "+7 900 000-00-12", 3, 1, "2026-10-15T14:00"),
    ("Игорь", "+7 900 000-00-13", 2, 3, "2026-10-15T11:00"),
    ("Ольга", "+7 900 000-00-14", 4, 4, "2026-10-15T12:00"),
    ("Наталья", "+7 900 000-00-15", 1, 2, "2026-10-16T15:00"),
    ("Павел", "+7 900 000-00-16", 2, 3, "2026-10-16T18:00"),
]


CSS = """
* { box-sizing: border-box; margin: 0; }
body { font: 15px/1.38 Inter, 'Noto Color Emoji', sans-serif; background: #fff; }
table { border-collapse: collapse; font: 12px/1.3 Inter, sans-serif; margin: 12px; }
td, th { border: 1px solid #d4d4d4; padding: 5px 8px; white-space: nowrap; }
th { background: #7b5ea7; color: #fff; font-weight: 600; }
tr:nth-child(odd) td { background: #f7f5fb; }
.sheet { padding: 0 0 4px; background: #fff; display: inline-block; }
.sheet h1 { font: 600 14px Inter, sans-serif; padding: 10px 12px 0; color: #333; }
"""


def render_sheet(xlsx: bytes) -> str:
    ws = load_workbook(io.BytesIO(xlsx)).active
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        tag = "th" if i == 0 else "td"
        cells = [
            v.strftime("%d.%m.%Y")
            if isinstance(v, datetime)
            else f"{v:,}".replace(",", " ")
            if isinstance(v, int) and v > 999
            else v
            for v in row
        ]
        rows.append("<tr>" + "".join(f"<{tag}>{escape(str(v))}</{tag}>" for v in cells) + "</tr>")
    return (
        f"<!doctype html><meta charset=utf-8><style>{CSS}</style>"
        f'<div class="sheet"><h1>Выгрузка для владельца — лист «{ws.title}»</h1>'
        f"<table>{''.join(rows)}</table></div>"
    )


async def main() -> None:
    catalog = load_catalog(ROOT / "data" / "catalog.json")
    storage = Storage(":memory:")
    core = BookingBot(catalog, storage, admins={":".join(OWNER)}, clock=lambda: datetime(2026, 10, 14, 11, 2))
    for name, phone, service_id, master_id, start in FAKE_CLIENTS:
        begin = datetime.fromisoformat(start)
        storage.book(
            platform="max" if master_id % 2 else "tg",
            user_id=f"9{master_id}",
            client_name=name,
            phone=phone,
            service_id=service_id,
            master_id=master_id,
            start=begin,
            end=begin + timedelta(minutes=catalog.services[service_id].duration_min),
            created_at=datetime(2026, 10, 12, 18, 30),
        )
    export = await core.handle_callback(*OWNER, "x")
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "tablitsa.html").write_text(render_sheet(export[0].file_bytes), encoding="utf-8")
    print(f"html: {OUT}")


if __name__ == "__main__":
    asyncio.run(main())
